from __future__ import annotations

import csv
import gzip
import re
from pathlib import Path
from typing import Iterable


ID_COLUMNS = {
    'id', 'sequenceid', 'sequence_id', 'sequence id', 'seqid', 'seq_id',
    'queryid', 'query_id', 'fastaid', 'fasta_id', 'isolateid',
    'isolate_id', 'isolate id', 'sampleid', 'sample_id',
}
PARTNER_ID_COLUMNS = {
    'partnerid', 'partner_id', 'partner id', 'partner', 'partneracronym',
    'partner_acronym', 'partner acronym',
}
GENOME_AVAILABLE_COLUMNS = {
    'alreadysequenced', 'already_sequenced', 'already sequenced',
    'genomeavailable', 'genome_available', 'genome available',
    'genomesequenced', 'genome_sequenced', 'genome sequenced',
    'wgsavailable', 'wgs_available', 'wgs available',
}
SELECTION_COLUMNS = {
    'selected', 'selectedforwgs', 'selected_for_wgs', 'selected for wgs',
    'selectedforgenomesequencing', 'selected_for_genome_sequencing',
    'selected for genome sequencing', 'selectedforfullgenomesequencing',
    'selected_for_full_genome_sequencing',
    'selected for full genome sequencing', 'wgs', 'wgsselected',
    'wgs_selected', 'wgs selected', 'genomesequencing',
    'genomesequencingselected', 'genome_sequencing_selected',
    'genome sequencing selected',
}
ADDITION_METHOD_COLUMNS = {
    'addition_method', 'isolate_addition_method', 'identification_method',
    'taxonomic_identification_method', 'submission_method',
}
_ADDITION_METHOD_ALIASES = {
    'marker': 'marker_sequence',
    'sequence': 'marker_sequence',
    'marker_sequence': 'marker_sequence',
    'marker_gene_sequence': 'marker_sequence',
    '16s': 'marker_sequence',
    '16s_sequence': 'marker_sequence',
    'maldi': 'maldi_tof',
    'maldi_tof': 'maldi_tof',
    'maldi_toff': 'maldi_tof',
    'marker_and_maldi': 'both',
    'marker_sequence_and_maldi_tof': 'both',
    'marker_gene_sequence_and_maldi_tof': 'both',
    'both': 'both',
}


def _normalise_column(name: object) -> str:
    return str(name or '').strip().lower().replace('-', '_')


def _normalise_for_matching(name: object) -> str:
    return _normalise_column(name).replace('_', '').replace(' ', '')


def _open_text(path: str | Path):
    if str(path).endswith('.gz'):
        return gzip.open(path, 'rt', newline='')
    return open(path, 'rt', newline='')


def _parse_boolean(value: object, *, field: str) -> bool:
    if value is None:
        return False
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return float(value) != 0.0
    text = str(value).strip().lower()
    if text in ('0', 'false', 'f', 'no', 'n'):
        return False
    if text in ('1', 'true', 't', 'yes', 'y'):
        return True
    raise ValueError(f'{field} must be an explicit yes/no Boolean; received {value!r}')


def _find_column(fieldnames: Iterable[str], candidates: set[str]) -> str | None:
    fields = list(fieldnames or [])
    exact = {_normalise_column(c): c for c in candidates}
    compact = {_normalise_for_matching(c): c for c in candidates}
    for field in fields:
        norm = _normalise_column(field)
        if norm in exact:
            return field
    for field in fields:
        norm = _normalise_for_matching(field)
        if norm in compact:
            return field
    return None


def _rows_from_delimited(path: str | Path) -> list[dict]:
    suffixes = ''.join(Path(str(path).removesuffix('.gz')).suffixes).lower()
    delimiter = '\t' if suffixes.endswith('.tsv') or suffixes.endswith('.txt') else ','
    with _open_text(path) as handle:
        sample = handle.read(4096)
        handle.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=',\t;')
            delimiter = dialect.delimiter
        except Exception:
            pass
        reader = csv.DictReader(handle, delimiter=delimiter)
        return [row for row in reader if any(v not in (None, '') for v in row.values())]


def load_partner_workbook_rows(
    path: str | Path,
    *,
    sheet_name: str = 'Bulk Upload Template',
) -> list[dict]:
    """Read the Rumen Gateway partner workbook into source-labeled isolate rows.

    This deliberately reads values only and never executes workbook macros. The
    source's ``taxon_id`` field is retained as a raw MALDI result because partner
    workbooks may contain a taxon name or a partner culture ID in that column.
    """
    try:
        from openpyxl import load_workbook
    except Exception as exc:
        raise RuntimeError('Reading .xlsm partner workbooks requires openpyxl') from exc

    workbook = load_workbook(path, read_only=True, data_only=True, keep_vba=True)
    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(
                f"Partner workbook does not contain sheet {sheet_name!r}; "
                f"available sheets: {', '.join(workbook.sheetnames)}"
            )
        sheet = workbook[sheet_name]
        rows_iter = sheet.iter_rows(values_only=True)
        header_row = next(rows_iter, None)
        if not header_row:
            return []
        headers = [str(value or '').strip().lower() for value in header_row]
        if 'record_id' not in headers:
            raise ValueError(f"Sheet {sheet_name!r} must have a record_id field in row 1")
        indices = {name: index for index, name in enumerate(headers) if name}

        def get(values, field):
            index = indices.get(field)
            return values[index] if index is not None and index < len(values) else None

        def clean(value):
            return str(value).strip() if value is not None else ''

        def enum_label(value):
            text = clean(value)
            return re.sub(r'^\s*\d+\s*,\s*', '', text).strip()

        result = []
        seen_ids = {}
        for row_number, values in enumerate(rows_iter, start=2):
            source_id = clean(get(values, 'record_id'))
            # The template has instructional rows between its internal header
            # and its first isolate. Rumen Gateway IDs use the stable PREFIX_NNNN form.
            if not re.fullmatch(r'[A-Za-z][A-Za-z0-9.-]*_\d+', source_id):
                continue
            if source_id in seen_ids:
                raise ValueError(
                    f'Partner workbook contains duplicate record_id {source_id!r} '
                    f'on rows {seen_ids[source_id]} and {row_number}'
                )
            seen_ids[source_id] = row_number
            partner_raw = clean(get(values, 'partner_id'))
            acronym_match = re.search(r'\(([A-Za-z0-9_-]+)\)\s*$', partner_raw)
            partner_id = acronym_match.group(1) if acronym_match else source_id.split('_', 1)[0]
            amplicon_method = enum_label(get(values, 'taxonomic_classification___1')).lower() == 'selected'
            maldi_method = enum_label(get(values, 'taxonomic_classification___2')).lower() == 'selected'
            maldi_result = clean(get(values, 'taxon_id'))

            metadata_map = {
                'partner_culture_id': 'partner_culture_id',
                'project': 'project',
                'microorganism_type': 'microorganism',
                'amplicon_method_reported': 'taxonomic_classification___1',
                'maldi_method_reported': 'taxonomic_classification___2',
                'amplicon_ncbi_name_reported': 'isolate_sci_name',
                'amplicon_ncbi_taxid_reported': 'isolate_ncbi_tax_id',
                'amplicon_ncbi_identity_reported': 'perc_identity_ncbi',
                'amplicon_gtdb_name_reported': 'gtdb_sci_name',
                'amplicon_gtdb_identity_reported': 'perc_identity_gtdb',
                'amplicon_gg2_name_reported': 'gg2_sci_name',
                'amplicon_gg2_identity_reported': 'perc_identity_gg2',
                'amplicon_silva_name_reported': 'silva_sci_name',
                'amplicon_silva_identity_reported': 'perc_identity_silva',
                'maldi_taxon_raw': 'taxon_id',
                'maldi_instrument': 'maldi_platform',
                'maldi_library': 'maldi_db_version',
                'host_species': 'ruminant_name',
                'host_taxid': 'ruminant_tax_id',
                'host_sex': 'ruminant_sex',
                'host_age_weeks': 'ruminant_age',
                'host_breed': 'ruminant_breed',
                'production_type': 'production_type',
                'sample_location': 'sample_location',
                'sample_description': 'sample_desc',
                'ruminant_diet': 'ruminant_diet',
                'diet_details': 'diet_details',
                'isolation_date': 'isolation_date',
                'isolation_temperature': 'isolation_temperature',
                'oxygen_requirement': 'oxygen_requirement',
                'culture_purity': 'pure_culture',
                'gram_stain': 'gram_stain',
                'culture_collection': 'deposite_culture_collect',
                'culture_collection_name': 'which_culture',
                'deposited_id': 'deposited_id',
            }
            metadata = {}
            for output_key, source_key in metadata_map.items():
                value = get(values, source_key)
                if value is None or not clean(value):
                    continue
                text = enum_label(value) if source_key in {
                    'taxonomic_classification___1', 'taxonomic_classification___2',
                    'microorganism', 'ruminant_name', 'ruminant_sex', 'ruminant_breed',
                    'production_type', 'sample_location', 'oxygen_requirement',
                    'pure_culture', 'gram_stain', 'deposite_culture_collect',
                } else clean(value)
                if text:
                    metadata[output_key] = text

            result.append({
                'source_row': row_number,
                'source_id': source_id,
                'partner_id': partner_id,
                'partner_raw': partner_raw,
                'amplicon_method_reported': amplicon_method,
                'maldi_method_reported': maldi_method,
                'maldi_result_raw': maldi_result,
                'maldi_instrument': clean(get(values, 'maldi_platform')),
                'maldi_library': clean(get(values, 'maldi_db_version')),
                'project_metadata': metadata,
            })
        return result
    finally:
        workbook.close()


def load_partner_sequencing_metadata(path: str | Path, sheet_name: str | None = None) -> list[dict]:
    """Read partner genome-availability metadata from CSV/TSV sidecar files.

    Returns rows with normalised keys:
      source_id, partner_id, selected_for_sequencing, genome_available,
      selected_for_wgs (stored alias for genome_available)
    """
    p = Path(path)
    lower = str(p).lower()
    if lower.endswith(('.csv', '.tsv', '.txt', '.csv.gz', '.tsv.gz', '.txt.gz')):
        raw_rows = _rows_from_delimited(p)
    else:
        raise ValueError(f"Unsupported partner metadata format: {path}. Use CSV/TSV sidecar metadata.")

    if not raw_rows:
        return []

    fieldnames = list(raw_rows[0].keys())
    id_col = _find_column(fieldnames, ID_COLUMNS)
    genome_col = _find_column(fieldnames, GENOME_AVAILABLE_COLUMNS)
    selection_col = _find_column(fieldnames, SELECTION_COLUMNS)
    partner_col = _find_column(fieldnames, PARTNER_ID_COLUMNS)
    addition_method_col = _find_column(fieldnames, ADDITION_METHOD_COLUMNS)

    if not id_col:
        raise ValueError(
            'Partner metadata must contain a sequence ID column matching FASTA IDs, such as sequence_id, isolate_id, sample_id, or ID.'
        )
    if not partner_col:
        raise ValueError(
            'Partner metadata must contain a partner acronym column such as partner_id, partner, or partner_acronym.'
        )
    if not genome_col:
        raise ValueError(
            'Partner metadata must contain an already-sequenced/genome-available column such as already_sequenced.'
        )

    rows = []
    seen: dict[str, tuple[str, bool, bool, int]] = {}
    for row_number, raw in enumerate(raw_rows, start=2):
        source_id = str(raw.get(id_col) or '').strip()
        if not source_id:
            continue
        partner_id = str(raw.get(partner_col) or '').strip()
        if not partner_id:
            continue
        raw_genome = raw.get(genome_col)
        raw_selection = raw.get(selection_col) if selection_col else 'no'
        genome_available = _parse_boolean(raw_genome, field='already_sequenced')
        selected_for_sequencing = _parse_boolean(
            raw_selection if raw_selection not in (None, '') else 'no',
            field='selected_for_genome_sequencing',
        )
        raw_addition_method = str(raw.get(addition_method_col) or '').strip() if addition_method_col else ''
        addition_method_key = _normalise_column(raw_addition_method).replace(' ', '_').replace('/', '_')
        addition_method = _ADDITION_METHOD_ALIASES.get(addition_method_key) if addition_method_key else None
        if addition_method_key and addition_method is None:
            raise ValueError(
                f'addition_method on row {row_number} must be marker_sequence, maldi_tof, or both; '
                f'received {raw_addition_method!r}'
            )
        signature = (
            partner_id.casefold(), selected_for_sequencing, genome_available, row_number,
        )
        if source_id in seen:
            previous = seen[source_id]
            duplicate_kind = (
                'duplicate' if previous[:3] == signature[:3] else 'conflicting'
            )
            raise ValueError(
                f'Partner metadata contains a {duplicate_kind} sequence_id {source_id!r} '
                f'on rows {previous[3]} and {row_number}; the cumulative ledger requires '
                'exactly one row per isolate.'
            )
        seen[source_id] = signature
        rows.append({
            'source_id': source_id,
            'partner_id': partner_id or source_id,
            'selected_for_sequencing': selected_for_sequencing,
            'genome_available': genome_available,
            # Existing project databases store genome availability in this
            # field; it does not mean current selection for sequencing.
            'selected_for_wgs': genome_available,
            'addition_method': addition_method,
            'raw_selected_value': '' if raw_genome is None else str(raw_genome),
            'raw_commitment_value': '' if raw_selection is None else str(raw_selection),
            'project_metadata': {
                str(key): str(value or '').strip()
                for key, value in raw.items()
                if key not in {id_col, partner_col, genome_col, selection_col, addition_method_col}
                and value not in (None, '')
            },
        })
    return rows
