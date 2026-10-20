"""Partner-submission validation before raw traces or project state are touched."""

from __future__ import annotations

import csv
import json
import re
from pathlib import Path
from typing import Iterable, Optional

from branchmanager.pipeline.paper_trail import (
    SEQUENCE_ID_COLUMNS,
    SUPPORTED_READ_EXTENSIONS,
    _metadata_file_columns,
    _metadata_key_variants,
    _is_supported_read_file,
    _normalise_row,
    _open_text,
    _row_processing_mode,
    _row_value,
)
from branchmanager.utils.fasta import read_fasta, write_fasta


PARTNER_COLUMNS = ('partner_id', 'partner', 'partner_acronym', 'partner_code')
GENOME_COLUMNS = ('already_sequenced', 'genome_available', 'genome_sequenced', 'wgs_available')
SELECTION_COLUMNS = (
    'selected_for_genome_sequencing', 'selected_for_wgs', 'wgs_selected',
    'genome_sequencing_selected',
)
MIXS_RECOMMENDED = (
    'sample_name', 'collection_date', 'geographic_location', 'host',
    'environment_broad_scale', 'environment_local_scale', 'environment_medium',
    'target_gene', 'pcr_primers', 'sequencing_method',
)
TRUE_VALUES = {'1', 'true', 'yes', 'y'}
FALSE_VALUES = {'0', 'false', 'no', 'n'}


def _read_xlsx_table(path: str | Path) -> list[dict[str, str]]:
    """Read the data sheet of an .xlsx file; unpivot 'Filename of primer sequence (X)' layouts."""
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise SystemExit('Reading .xlsx metadata requires openpyxl (pip install openpyxl).') from exc
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheets = [list(sheet.iter_rows(values_only=True)) for sheet in workbook.worksheets]
    finally:
        workbook.close()

    def cleaned(rows):
        return [r for r in rows if any(c is not None and str(c).strip() for c in r)]

    def is_data(rows):
        return len(rows) > 1 and sum(c is not None for c in rows[0]) > 1

    candidates = [cleaned(rows) for rows in sheets]
    candidates = [rows for rows in candidates if is_data(rows)]
    if not candidates:
        return []
    rows = candidates[0]
    header = [' '.join(str(c).split()) if c is not None else '' for c in rows[0]]
    primer_re = re.compile(r'^filename of primer sequence \(([^)]+)\)$', re.I)
    primer_cols = {i: m.group(1).strip() for i, h in enumerate(header) if (m := primer_re.match(h))}
    if primer_cols:
        out = []
        for row in rows[1:]:
            isolate = '' if row[0] is None else str(row[0]).strip()
            if not isolate or isolate.lower().startswith('ensure that'):
                continue
            for i, primer in primer_cols.items():
                value = row[i] if i < len(row) else None
                if value is not None and str(value).strip():
                    out.append(_normalise_row({
                        'sequence_id': isolate,
                        'supplier_read_id': str(value).strip(),
                        'primer': primer,
                    }))
        return out
    return [
        _normalise_row({
            name: ('' if cell is None else str(cell))
            for name, cell in zip(header, row) if name
        })
        for row in rows[1:]
    ]


def _read_table(path: str | Path) -> list[dict[str, str]]:
    if str(path).lower().endswith(('.xlsx', '.xlsm')):
        return _read_xlsx_table(path)
    with _open_text(path) as handle:
        sample = handle.read(4096)
        handle.seek(0)
        delimiter = '\t' if str(path).lower().endswith(('.tsv', '.tsv.gz', '.txt', '.txt.gz')) else ','
        try:
            delimiter = csv.Sniffer().sniff(sample, delimiters=',\t').delimiter
        except Exception:
            pass
        return [_normalise_row(row) for row in csv.DictReader(handle, delimiter=delimiter)]


def _populated_fields(
    rows: list[dict],
    candidates: Iterable[str],
    *,
    required: Iterable[str] = (),
) -> list[str]:
    """Keep required fields and optional fields containing at least one value."""
    required_fields = set(required)
    return [
        field for field in candidates
        if field in required_fields
        or any(str(row.get(field) or '').strip() for row in rows)
    ]


def _resolve_file_matches(
    value: str,
    table_path: str | Path,
    read_dir: Optional[str | Path],
) -> tuple[list[Path], Path]:
    item = Path(value).expanduser()
    if item.is_absolute():
        return ([item] if item.is_file() else []), item
    candidates = [Path(table_path).resolve().parent / item]
    if read_dir:
        candidates.insert(0, Path(read_dir).expanduser().resolve() / item)
    for candidate in candidates:
        if candidate.is_file():
            return [candidate], candidate

    requested = set(_metadata_key_variants(value))
    matches = []
    search_roots = []
    if read_dir:
        search_roots.append(Path(read_dir).expanduser().resolve())
    search_roots.append(Path(table_path).resolve().parent)
    seen = set()
    for root in search_roots:
        if not root.is_dir():
            continue
        for child in root.rglob('*'):
            if not child.is_file() or not _is_supported_read_file(child):
                continue
            child_keys = set(_metadata_key_variants(child))
            child_keys.add(child.stem.split('_', 1)[0])
            if not requested.intersection(child_keys):
                continue
            resolved = child.resolve()
            if str(resolved) not in seen:
                matches.append(resolved)
                seen.add(str(resolved))
    return sorted(matches), candidates[0]


def validate_submission(
    sample_map: Optional[str | Path] = None,
    *,
    fasta: Optional[str | Path] = None,
    partner_metadata: Optional[str | Path] = None,
    read_dir: Optional[str | Path] = None,
    primers: Iterable[str] = (),
    expected_partner_id: Optional[str] = None,
    dataset: Optional[str] = None,
) -> dict:
    """Validate IDs, raw-file mapping, metadata coverage, and duplicate ownership."""
    if bool(sample_map) == bool(fasta):
        raise ValueError('provide exactly one of sample_map or fasta')
    input_type = 'ab1_map' if sample_map else 'fasta'
    rows = _read_table(sample_map) if sample_map else []
    problems = []
    normalised_reads = []
    seen_ids = set()
    normalised_by_id = {}
    seen_read_assignments = set()
    file_owner = {}
    partner_by_id = {}
    genome_by_id = {}
    selection_by_id = {}
    metadata_line_by_id = {}
    metadata_signature_by_id = {}
    expected_partner = str(expected_partner_id or '').strip()
    submission_dataset = str(dataset or '').strip()

    metadata_path = partner_metadata
    if metadata_path is None:
        raise ValueError('partner_metadata is required for every submission')
    metadata_rows = _read_table(metadata_path)
    for line_number, row in enumerate(metadata_rows, start=2):
        sequence_id = _row_value(row, *SEQUENCE_ID_COLUMNS)
        if not sequence_id:
            continue
        partner = _row_value(row, *PARTNER_COLUMNS)
        raw_genome = _row_value(row, *GENOME_COLUMNS)
        raw_selection = _row_value(row, *SELECTION_COLUMNS)
        signature = (partner.casefold(), raw_genome.casefold(), raw_selection.casefold())
        if sequence_id in metadata_signature_by_id:
            first_line = metadata_line_by_id[sequence_id]
            code = (
                'DUPLICATE_METADATA_ID'
                if metadata_signature_by_id[sequence_id] == signature
                else 'CONFLICTING_METADATA_ID'
            )
            problems.append({
                'severity': 'ERROR', 'line': line_number, 'sequence_id': sequence_id,
                'code': code,
                'detail': f'sequence_id already appears on metadata line {first_line}; the cumulative ledger requires one row per isolate',
            })
            continue
        metadata_line_by_id[sequence_id] = line_number
        metadata_signature_by_id[sequence_id] = signature
        partner_by_id[sequence_id] = partner
        genome_by_id[sequence_id] = raw_genome
        selection_by_id[sequence_id] = raw_selection
        if not partner:
            problems.append({'severity': 'ERROR', 'line': line_number, 'sequence_id': sequence_id, 'code': 'MISSING_PARTNER_ID', 'detail': 'partner acronym is required'})
        if not raw_genome:
            problems.append({'severity': 'ERROR', 'line': line_number, 'sequence_id': sequence_id, 'code': 'MISSING_GENOME_STATUS', 'detail': 'already_sequenced/genome_available must be explicitly yes or no'})
        elif raw_genome.strip().lower() not in TRUE_VALUES | FALSE_VALUES:
            problems.append({'severity': 'ERROR', 'line': line_number, 'sequence_id': sequence_id, 'code': 'INVALID_GENOME_STATUS', 'detail': f'unrecognised Boolean value {raw_genome!r}'})
        if raw_selection and raw_selection.strip().lower() not in TRUE_VALUES | FALSE_VALUES:
            problems.append({'severity': 'ERROR', 'line': line_number, 'sequence_id': sequence_id, 'code': 'INVALID_SELECTION_STATUS', 'detail': f'unrecognised Boolean value {raw_selection!r}'})

    if fasta:
        fasta_path = Path(fasta).expanduser().resolve()
        if not fasta_path.is_file():
            problems.append({'severity': 'ERROR', 'line': '', 'sequence_id': '', 'code': 'FASTA_NOT_FOUND', 'detail': str(fasta_path)})
            fasta_records = []
        else:
            fasta_records = list(read_fasta(str(fasta_path)))
        if not fasta_records:
            problems.append({'severity': 'ERROR', 'line': '', 'sequence_id': '', 'code': 'EMPTY_FASTA', 'detail': str(fasta_path)})
        normalised_fasta_records = []
        for record_number, (sequence_id, sequence) in enumerate(fasta_records, start=1):
            sequence_id = str(sequence_id).strip()
            if not sequence_id:
                problems.append({'severity': 'ERROR', 'line': record_number, 'sequence_id': '', 'code': 'MISSING_SEQUENCE_ID', 'detail': 'FASTA header is empty'})
                continue
            if sequence_id in seen_ids:
                problems.append({'severity': 'ERROR', 'line': record_number, 'sequence_id': sequence_id, 'code': 'DUPLICATE_SEQUENCE_ID', 'detail': 'FASTA sequence IDs must be unique across the submission'})
                continue
            seen_ids.add(sequence_id)
            cleaned = ''.join(str(sequence or '').split()).upper().replace('U', 'T')
            if not cleaned:
                problems.append({'severity': 'ERROR', 'line': record_number, 'sequence_id': sequence_id, 'code': 'EMPTY_SEQUENCE', 'detail': 'FASTA record has no sequence'})
            invalid = sorted(set(cleaned) - set('ACGTRYSWKMBDHVN.-'))
            if invalid:
                problems.append({'severity': 'ERROR', 'line': record_number, 'sequence_id': sequence_id, 'code': 'INVALID_SEQUENCE_CHARACTERS', 'detail': ''.join(invalid)})
            entry = {
                'sequence_id': sequence_id,
                'partner_id': partner_by_id.get(sequence_id, ''),
                'dataset': submission_dataset,
                'selected_for_genome_sequencing': selection_by_id.get(sequence_id, '') or 'no',
                'already_sequenced': genome_by_id.get(sequence_id, ''),
                'processing_mode': 'provided_fasta',
                'input_type': 'fasta',
                'source_fasta': str(fasta_path),
                'read_files': '',
                'primers_from_columns': '',
                **{column: '' for column in MIXS_RECOMMENDED},
            }
            normalised_by_id[sequence_id] = entry
            normalised_fasta_records.append((sequence_id, cleaned))
    else:
        normalised_fasta_records = []

    map_partner_ids = set()
    map_datasets = set()
    for line_number, row in enumerate(rows, start=2):
        sequence_id = _row_value(row, *SEQUENCE_ID_COLUMNS)
        if not sequence_id:
            problems.append({'severity': 'ERROR', 'line': line_number, 'sequence_id': '', 'code': 'MISSING_SEQUENCE_ID', 'detail': 'sequence/isolate ID is required'})
            continue
        seen_ids.add(sequence_id)
        mode = _row_processing_mode(row) or 'assemble'
        row_partner = _row_value(row, *PARTNER_COLUMNS)
        row_dataset = _row_value(row, 'dataset', 'batch_id', 'batch', 'submission_id')
        if row_partner:
            map_partner_ids.add(row_partner)
        if row_dataset:
            map_datasets.add(row_dataset)
        entry = normalised_by_id.setdefault(sequence_id, {
            'sequence_id': sequence_id,
            'partner_id': row_partner or partner_by_id.get(sequence_id, '') or expected_partner,
            'dataset': row_dataset or submission_dataset,
            'selected_for_genome_sequencing': selection_by_id.get(sequence_id, '') or 'no',
            'already_sequenced': genome_by_id.get(sequence_id, ''),
            'processing_mode': mode,
            'input_type': 'ab1_map',
            'source_fasta': '',
            '_read_files': [],
            '_primers': [],
            '_line': line_number,
            **{column: '' for column in MIXS_RECOMMENDED},
        })
        if entry['processing_mode'] != mode:
            problems.append({'severity': 'ERROR', 'line': line_number, 'sequence_id': sequence_id, 'code': 'CONFLICTING_PROCESSING_MODE', 'detail': f"{entry['processing_mode']} versus {mode}"})
        if row_partner and entry['partner_id'] and row_partner.casefold() != str(entry['partner_id']).casefold():
            problems.append({'severity': 'ERROR', 'line': line_number, 'sequence_id': sequence_id, 'code': 'CONFLICTING_PARTNER_ID', 'detail': f"{entry['partner_id']} versus {row_partner}"})
        if row_dataset and entry['dataset'] and row_dataset.casefold() != str(entry['dataset']).casefold():
            problems.append({'severity': 'ERROR', 'line': line_number, 'sequence_id': sequence_id, 'code': 'CONFLICTING_DATASET', 'detail': f"{entry['dataset']} versus {row_dataset}"})
        for column in MIXS_RECOMMENDED:
            if row.get(column) and not entry.get(column):
                entry[column] = row[column]

        file_cells = _metadata_file_columns(row, primers)
        if not file_cells:
            problems.append({'severity': 'ERROR', 'line': line_number, 'sequence_id': sequence_id, 'code': 'NO_READ_FILES', 'detail': 'no AB1/ABI/FASTA/FASTQ files were mapped'})
            continue
        row_primer = _row_value(row, 'primer', 'primer_name')
        row_direction = _row_value(row, 'direction', 'orientation', 'read', 'read_direction').lower()
        for value, primer_from_column in file_cells:
            matches, fallback = _resolve_file_matches(value, sample_map, read_dir)
            if not matches:
                problems.append({'severity': 'ERROR', 'line': line_number, 'sequence_id': sequence_id, 'code': 'READ_FILE_NOT_FOUND', 'detail': str(fallback)})
                continue
            if len(matches) > 1:
                problems.append({'severity': 'ERROR', 'line': line_number, 'sequence_id': sequence_id, 'code': 'AMBIGUOUS_READ_FILE', 'detail': ';'.join(str(path) for path in matches)})
                continue
            resolved = matches[0]
            if not str(resolved).lower().endswith(SUPPORTED_READ_EXTENSIONS):
                problems.append({'severity': 'ERROR', 'line': line_number, 'sequence_id': sequence_id, 'code': 'UNSUPPORTED_READ_FILE', 'detail': str(resolved)})
                continue
            owner = file_owner.get(str(resolved.resolve()))
            if owner and owner != sequence_id:
                problems.append({'severity': 'ERROR', 'line': line_number, 'sequence_id': sequence_id, 'code': 'READ_FILE_MULTIPLE_ISOLATES', 'detail': f'{resolved} is also assigned to {owner}'})
            file_owner[str(resolved.resolve())] = sequence_id
            assignment = (sequence_id, str(resolved.resolve()))
            if assignment in seen_read_assignments:
                continue
            seen_read_assignments.add(assignment)
            primer = row_primer or primer_from_column
            direction = row_direction if row_direction in {'forward', 'reverse'} else (
                'reverse' if str(primer).upper().endswith('R') else ('forward' if primer else '')
            )
            entry['_read_files'].append(str(resolved.resolve()))
            entry['_primers'].append(primer)
            normalised_reads.append({
                'sequence_id': sequence_id,
                'partner_id': entry['partner_id'],
                'dataset': entry['dataset'],
                'file': str(resolved.resolve()),
                'primer': primer,
                'direction': direction,
                'processing_mode': entry['processing_mode'],
            })

    normalised = []
    for sequence_id, entry in normalised_by_id.items():
        line_number = entry.pop('_line', '')
        if sequence_id not in partner_by_id:
            problems.append({'severity': 'ERROR', 'line': line_number, 'sequence_id': sequence_id, 'code': 'MISSING_PARTNER_METADATA_ROW', 'detail': f'no matching row in {metadata_path}'})
        metadata_partner = partner_by_id.get(sequence_id, '')
        if expected_partner and metadata_partner and metadata_partner.casefold() != expected_partner.casefold():
            problems.append({
                'severity': 'ERROR', 'line': line_number, 'sequence_id': sequence_id,
                'code': 'PARTNER_ID_MISMATCH',
                'detail': f'ledger partner {metadata_partner!r} does not match expected partner {expected_partner!r}',
            })
        if expected_partner and entry.get('partner_id') and str(entry['partner_id']).casefold() != expected_partner.casefold():
            problems.append({
                'severity': 'ERROR', 'line': line_number, 'sequence_id': sequence_id,
                'code': 'PARTNER_ID_MISMATCH',
                'detail': f'sample-map partner {entry["partner_id"]!r} does not match expected partner {expected_partner!r}',
            })
        if expected_partner and not entry.get('partner_id'):
            entry['partner_id'] = expected_partner
        if submission_dataset and not entry.get('dataset'):
            entry['dataset'] = submission_dataset
        if '_read_files' in entry:
            entry['read_files'] = ';'.join(entry.pop('_read_files'))
            entry['primers_from_columns'] = ';'.join(primer for primer in entry.pop('_primers') if primer)
        normalised.append(entry)

    if len(map_partner_ids) > 1:
        problems.append({'severity': 'ERROR', 'line': '', 'sequence_id': '', 'code': 'MULTIPLE_PARTNERS_IN_SAMPLE_MAP', 'detail': ','.join(sorted(map_partner_ids))})
    if len(map_datasets) > 1:
        problems.append({'severity': 'ERROR', 'line': '', 'sequence_id': '', 'code': 'MULTIPLE_DATASETS_IN_SAMPLE_MAP', 'detail': ','.join(sorted(map_datasets))})

    if read_dir and Path(read_dir).expanduser().is_dir():
        mapped_files = set(file_owner)
        physical_files = sorted(
            child.resolve()
            for child in Path(read_dir).expanduser().resolve().rglob('*')
            if child.is_file() and _is_supported_read_file(child)
        )
        for path in physical_files:
            if str(path) not in mapped_files:
                problems.append({
                    'severity': 'WARNING', 'line': '', 'sequence_id': '',
                    'code': 'UNMAPPED_READ_FILE', 'detail': str(path),
                })
    error_count = sum(row['severity'] == 'ERROR' for row in problems)
    return {
        'status': 'PASS' if error_count == 0 else 'FAIL',
        'input_type': input_type,
        'isolates': len(normalised),
        'read_files': len(file_owner),
        'errors': error_count,
        'warnings': sum(row['severity'] == 'WARNING' for row in problems),
        'normalised': normalised,
        'normalised_reads': normalised_reads,
        'normalised_fasta_records': normalised_fasta_records,
        'problems': problems,
    }


def write_onboarding_outputs(outdir: str | Path, result: dict) -> dict:
    output = Path(outdir)
    output.mkdir(parents=True, exist_ok=True)
    normalised_path = output / 'normalised_submission.tsv'
    normalised_candidates = [
        'sequence_id', 'partner_id', 'dataset', 'selected_for_genome_sequencing', 'already_sequenced',
        'input_type', 'processing_mode', 'source_fasta', 'read_files',
        'primers_from_columns', *MIXS_RECOMMENDED,
    ]
    normalised_fields = _populated_fields(
        result['normalised'],
        normalised_candidates,
        required=(
            'sequence_id', 'partner_id', 'dataset',
            'selected_for_genome_sequencing', 'already_sequenced',
            'input_type', 'processing_mode',
        ),
    )
    with open(normalised_path, 'w', newline='') as handle:
        writer = csv.DictWriter(handle, fieldnames=normalised_fields, delimiter='\t', extrasaction='ignore')
        writer.writeheader()
        writer.writerows(result['normalised'])
    report_path = output / 'onboarding_report.tsv'
    with open(report_path, 'w', newline='') as handle:
        fields = ['severity', 'line', 'sequence_id', 'code', 'detail']
        writer = csv.DictWriter(handle, fieldnames=fields, delimiter='\t', extrasaction='ignore')
        writer.writeheader()
        writer.writerows(
            row for row in result['problems']
            if row.get('severity') in {'ERROR', 'WARNING'}
        )
    summary_path = output / 'onboarding_summary.json'
    summary_path.write_text(json.dumps({key: result[key] for key in ('status', 'input_type', 'isolates', 'read_files', 'errors', 'warnings')}, indent=2) + '\n')
    read_map_path = output / 'normalised_read_map.tsv'
    with open(read_map_path, 'w', newline='') as handle:
        field_candidates = [
            'sequence_id', 'file', 'primer', 'direction', 'processing_mode',
            'partner_id', 'dataset',
        ]
        fields = _populated_fields(
            result.get('normalised_reads', []),
            field_candidates,
            required=(
                'sequence_id', 'file', 'direction', 'processing_mode',
                'partner_id', 'dataset',
            ),
        )
        writer = csv.DictWriter(
            handle, fieldnames=fields, delimiter='\t', extrasaction='ignore',
        )
        writer.writeheader()
        writer.writerows(result.get('normalised_reads', []))
    outputs = {
        'normalised': str(normalised_path), 'read_map': str(read_map_path),
        'report': str(report_path), 'summary': str(summary_path),
    }
    if result.get('input_type') == 'fasta':
        fasta_path = output / 'normalised_input.fasta'
        write_fasta(result.get('normalised_fasta_records', []), str(fasta_path))
        outputs['fasta'] = str(fasta_path)
    return outputs
